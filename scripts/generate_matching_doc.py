"""Regenerate ``docs/Heuristic_DeepLearning_Matching_Implementation.docx``.

Content is aligned with ``src/cross`` (Path B greedy/Hungarian, ranker MLP,
TemporalGraphRanker, calibration / evidence). Requires ``python-docx``.

If the target docx already exists, the first PNG embedded under ``word/media/``
is copied into the new file as the workflow figure (optional).
"""
from __future__ import annotations

import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "Heuristic_DeepLearning_Matching_Implementation.docx"


def _extract_first_media_png(docx_path: Path) -> Path | None:
    if not docx_path.is_file():
        return None
    try:
        with zipfile.ZipFile(docx_path) as zf:
            for name in sorted(zf.namelist()):
                if name.startswith("word/media/") and name.lower().endswith(".png"):
                    fd, tmp = tempfile.mkstemp(suffix=".png")
                    import os

                    os.close(fd)
                    outp = Path(tmp)
                    outp.write_bytes(zf.read(name))
                    return outp
    except OSError:
        return None
    return None


def _add_body(doc: Document) -> None:
    doc.add_heading("启发式与深度学习协同的跨链取款匹配（实现说明）", level=0)

    doc.add_heading("一、文档范围与代码映射", level=1)
    doc.add_paragraph(
        "本文描述 Celer Cross 仓库中 Path B 的「ETH 存款 → BNB 取款」匹配思路与实现要点，"
        "对应代码主要在：cross.domain.path_b（候选与时间窗）、cross.domain.ranker（RankerMLP）、"
        "cross.domain.graph（TemporalGraphRanker）、cross.domain.path_b.greedy（分配与证据）。"
        "入口与参数见仓库 README 与 cross.interfaces.cli。"
    )

    doc.add_heading("二、业务总览", level=1)
    doc.add_paragraph(
        "输入为 ETH 侧 Celer 存款类交易；经 AML 规则或模型过滤后得到较小的高可疑子集。"
        "对每笔待追踪的源交易：若其源交易哈希在标签文件 label 中命中，则直接用标签中的目标哈希（dst）"
        "经链上 API 拉取 BNB 交易并成对输出；若未命中标签，则在可调时间窗口内向链上拉取候选 BNB 取款交易，"
        "在候选集合上做启发式误差与可选学习式打分，再做全局一对一分配，输出与 AML 源交易对应的 BNB 取款交易。"
    )
    p = doc.add_paragraph()
    p.add_run("数据流摘要：").bold = True
    p.add_run(
        "AML 过滤后的子集 → 时间窗内 BNB 候选 →（可选）MLP / 候选注意力模型重标定边代价 → "
        "贪心或匈牙利二分匹配 → 输出对与证据字段。"
    )

    doc.add_heading("三、启发式与可学习打分", level=1)

    doc.add_heading("3.1 启发式（基误差）", level=2)
    doc.add_paragraph(
        "对每条 (源行, 候选 dst 哈希) 边，先由 compute_base_error_and_features 计算标量基误差 base_err "
        "（越小越好）及 9 维特征向量。基误差综合金额一致性（有代币汇率时按 raw 金额比；否则在人类可读金额上比较）"
        "与相对桥接延迟（相对历史中位延迟，带 delay_weight）。"
    )
    doc.add_paragraph("启发式侧用到的主要分量：", style="List Bullet")
    doc.add_paragraph("金额比率误差：在可得 ETH→BNB 代币汇率时，期望 raw 转出与观测 raw 的相对偏差。", style="List Bullet")
    doc.add_paragraph("人类可读相对误差：在双方人类可读金额可得时，|BNB−ETH_human| / ETH_human。", style="List Bullet")
    doc.add_paragraph("延迟误差：取款最早时间相对源时间戳的间隔，与中位桥接延迟的归一化偏差（并入总误差）。", style="List Bullet")

    doc.add_heading("3.2 RankerMLP（边级）", level=2)
    doc.add_paragraph(
        "类名 RankerMLP：对单条边使用与启发式相同的 9 维特征；推理时用训练 checkpoint 中的均值方差标准化，"
        "经小型 MLP 后 sigmoid 得到匹配概率 p，将边代价调整为 base_err × (2 − p)。概率越高，最终代价越低。"
    )
    doc.add_paragraph("9 维特征依次为：", style="List Bullet")
    doc.add_paragraph("log1p(ETH raw 金额)、log1p(BNB 侧 raw 合计)、比率误差 clip、人类可读相对误差 clip、延迟归一化项、"
                      "时间差 / 窗口归一化、比率是否可用标志、人类可读金额是否同时可用标志、综合基误差 clip。", style="List Bullet")

    doc.add_heading("3.3 TemporalGraphRanker（源内多候选）", level=2)
    doc.add_paragraph(
        "类名 TemporalGraphRanker：对同一源交易，在时间窗内取至多 top_k 个候选 dst，"
        "每条边使用 9 维 Path B 特征并拼接 6 维位置/排行尾部特征（共 GRAPH_EDGE_DIM 维），"
        "形成候选序列；经线性投影与 PyTorch MultiheadAttention（batch 维为 1 的「单源多候选」序列），"
        "再经前馈头输出 logits，对候选维做 softmax 得到竞争概率 p，边代价同样为 base_err × (2 − p)。"
        "这是「候选集上的自注意力排序」，与经典图神经网络在一般拓扑上的消息传递不同，但可表达候选间相对关系。"
    )

    doc.add_heading("3.4 hybrid 模式", level=2)
    doc.add_paragraph(
        "build_greedy_edge_score_fn 支持 ranker_mode=heuristic | mlp | graph | hybrid。"
        "hybrid 且两支 checkpoint 均可用时：对同一边分别计算 MLP 调整代价与 Graph 调整代价，"
        "再按 (hybrid_heuristic_weight × mlp_cost + hybrid_graph_weight × graph_cost) / (权重和) 凸组合（参数名沿历史命名，实为 MLP 与 Graph 两支）。"
        "若仅一支可用则回退为单支。"
    )

    doc.add_heading("四、候选与匹配", level=1)

    doc.add_heading("4.1 两阶段候选（first / second batch）", level=2)
    doc.add_paragraph(
        "greedy_pair_dst_hashes 内对每个源分两批截取候选：先按较严规则取 first_batch_k（默认 100），"
        "再放宽取 second_batch_k（默认 100），以平衡精度与召回；具体阈值与过滤见 greedy 模块内实现。"
    )

    doc.add_heading("4.2 二分图分配", level=2)
    doc.add_paragraph("匈牙利算法（linear_sum_assignment）：", style="List Bullet")
    doc.add_paragraph("在剩余源与候选 dst 并上每行 dummy 列上求最小代价完美匹配；理论最优（在给定离散候选与代价下）。", style="List Bullet")
    doc.add_paragraph("稠密代价矩阵约为 R×(M+R)×8 字节；超过约 400 MiB 或源行数 > 8000 时自动回退为全局贪心，避免内存与 O(R³) 风险。", style="List Bullet")
    doc.add_paragraph("贪心分配：全部 (代价, 源, dst) 边全局排序，先到先得；复杂度约 O(n log n)，适合大规模。", style="List Bullet")

    doc.add_heading("五、可解释性与输出证据", level=1)

    doc.add_heading("5.1 置信度分段校准", level=2)
    doc.add_paragraph(
        "_calibrate_confidence：对原始 softmax 置信度做分段线性压缩，高段上限约 0.93，"
        "避免数值上「过度自信」，便于风控复核。"
    )

    doc.add_heading("5.2 不确定性分层", level=2)
    doc.add_paragraph(
        "_uncertainty_band：结合候选数量与最优/次优代价间隙的相对值，输出 high / medium / low；"
        "候选≤1 或无二名时偏 high；相对间隙 ≥1.0 偏 low；≥0.3 为 medium；否则 high。"
    )

    doc.add_heading("5.3 证据字段（节选）", level=2)
    doc.add_paragraph(
        "每源可含：selected_dstTxHash、selected_rank、selected_confidence、selected_confidence_calibrated、"
        "uncertainty_band、counter_evidence、topk_candidates、叙事 narrative 等，供审计与人工复核。"
    )

    doc.add_heading("六、工作流示意图", level=1)
    doc.add_paragraph(
        "下图若存在，为上一版文档内嵌流程图；亦可对照仓库 README 中 Path B / 分层架构说明。"
    )


def main() -> None:
    media_tmp: Path | None = _extract_first_media_png(OUT_DOCX)

    doc = Document()
    _add_body(doc)

    if media_tmp is not None:
        try:
            doc.add_picture(str(media_tmp), width=Cm(14))
            last = doc.paragraphs[-1]
            last.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        finally:
            try:
                media_tmp.unlink(missing_ok=True)
            except OSError:
                pass

    doc.add_heading("七、设计要点小结", level=1)
    for text in (
        "两阶段候选截取，平衡效率与召回。",
        "时间窗与接收方解析（如 bnb_pick_per_candidate）随链上数据自适应。",
        "MLP 侧重单边特征；TemporalGraphRanker 侧重同源候选间竞争（softmax）。",
        "保守置信度校准与不确定性、反证字段，服务合规复核。",
        "匈牙利带内存/行数护栏，大规模自动降级贪心。",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("八、生成信息", level=1)
    gen = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    doc.add_paragraph(f"由 scripts/generate_matching_doc.py 生成。\n生成时间：{gen}\n输出文件：{OUT_DOCX}\n依赖：pip install python-docx")

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
